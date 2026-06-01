import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent))

import tempfile
from datetime import date
from decimal import Decimal

import openpyxl

from src.models.transaction import Transaction, TransactionType, Category, Invoice
from src.reports.excel_report import generate_report


def make_txn(d, desc, amount, ttype=TransactionType.DEBIT, cat=Category.EXPENSE):
    return Transaction(
        date=date.fromisoformat(d),
        description=desc,
        amount=Decimal(str(amount)),
        transaction_type=ttype,
        category=cat,
    )


def make_invoice(num, issue, due, vendor, amount, paid=False, payment_date=None):
    inv = Invoice(
        invoice_number=num,
        issue_date=date.fromisoformat(issue),
        due_date=date.fromisoformat(due),
        vendor=vendor,
        line_items=[{"description": "Service", "amount": str(amount)}],
        paid=paid,
        payment_date=date.fromisoformat(payment_date) if payment_date else None,
    )
    return inv


def test_creates_file():
    txns = [make_txn("2026-05-01", "Rent", 1000)]
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "report.xlsx"
        result = generate_report(txns, output_path=out)
        assert result.exists()
        assert result.suffix == ".xlsx"


def test_summary_sheet_present():
    txns = [make_txn("2026-05-01", "Rent", 1000)]
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "report.xlsx"
        generate_report(txns, output_path=out)
        wb = openpyxl.load_workbook(out)
        assert "Financial Summary" in wb.sheetnames


def test_transaction_detail_sheet_present():
    txns = [make_txn("2026-05-01", "Rent", 1000)]
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "report.xlsx"
        generate_report(txns, output_path=out)
        wb = openpyxl.load_workbook(out)
        assert "Transaction Detail" in wb.sheetnames


def test_ar_aging_sheet_present_when_invoices_given():
    txns = [make_txn("2026-05-01", "Rent", 1000)]
    invs = [make_invoice("INV-001", "2026-05-01", "2026-06-01", "Acme", 500)]
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "report.xlsx"
        generate_report(txns, invoices=invs, output_path=out)
        wb = openpyxl.load_workbook(out)
        assert "AR Aging" in wb.sheetnames


def test_ar_aging_sheet_absent_when_no_invoices():
    txns = [make_txn("2026-05-01", "Rent", 1000)]
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "report.xlsx"
        generate_report(txns, output_path=out)
        wb = openpyxl.load_workbook(out)
        assert "AR Aging" not in wb.sheetnames


def test_transaction_rows_match_count():
    txns = [
        make_txn("2026-05-01", "Rent", 1000),
        make_txn("2026-05-02", "Revenue", 2000, TransactionType.CREDIT, Category.REVENUE),
        make_txn("2026-05-03", "Supplies", 300),
    ]
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "report.xlsx"
        generate_report(txns, output_path=out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Transaction Detail"]
        # Row 1 = headers, rows 2..N = data
        data_rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]
        assert len(data_rows) == 3


def test_only_unpaid_invoices_in_aging():
    invs = [
        make_invoice("INV-001", "2026-05-01", "2026-05-15", "Acme", 500, paid=True, payment_date="2026-05-14"),
        make_invoice("INV-002", "2026-05-01", "2026-05-20", "Beta", 300, paid=False),
    ]
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "report.xlsx"
        generate_report([], invoices=invs, output_path=out, report_date=date(2026, 5, 31))
        wb = openpyxl.load_workbook(out)
        ws = wb["AR Aging"]
        # Row 3 = header, rows 4+ = data — check only 1 data row
        data_rows = [
            r for r in ws.iter_rows(min_row=4, max_row=10, values_only=True)
            if r[0] and str(r[0]).startswith("INV")
        ]
        assert len(data_rows) == 1
        assert data_rows[0][0] == "INV-002"


if __name__ == "__main__":
    tests = [
        test_creates_file,
        test_summary_sheet_present,
        test_transaction_detail_sheet_present,
        test_ar_aging_sheet_present_when_invoices_given,
        test_ar_aging_sheet_absent_when_no_invoices,
        test_transaction_rows_match_count,
        test_only_unpaid_invoices_in_aging,
    ]
    passed = 0
    for t in tests:
        try:
            t()
            print(f"  PASS  {t.__name__}")
            passed += 1
        except Exception as e:
            print(f"  FAIL  {t.__name__}: {e}")
    print(f"\n{passed}/{len(tests)} tests passed")
