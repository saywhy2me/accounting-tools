"""Excel report generator.

Produces a multi-sheet .xlsx workbook:
  Sheet 1 — Financial Summary   (income/expense totals by category, net position)
  Sheet 2 — Transaction Detail  (full transaction list with colour-coded types)
  Sheet 3 — AR Aging            (outstanding invoices bucketed by age)
"""

from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import List, Optional

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import (Alignment, Border, Font, PatternFill, Side,
                              numbers as xl_numbers)
from openpyxl.utils import get_column_letter

from src.models.transaction import Invoice, Transaction, TransactionType, Category
from src.utils.currency import format_currency


# ── palette ──────────────────────────────────────────────────────────────────
_HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")
_CREDIT_FILL   = PatternFill("solid", fgColor="E2EFDA")
_DEBIT_FILL    = PatternFill("solid", fgColor="FCE4D6")
_OVERDUE_FILL  = PatternFill("solid", fgColor="FF0000")
_WARNING_FILL  = PatternFill("solid", fgColor="FFEB9C")
_TOTAL_FILL    = PatternFill("solid", fgColor="D9E1F2")
_HEADER_FONT   = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
_BOLD_FONT     = Font(bold=True, name="Calibri")
_NORMAL_FONT   = Font(name="Calibri")
_THIN          = Side(style="thin")
_BORDER        = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CURRENCY_FMT  = '#,##0.00'
_DATE_FMT      = 'YYYY-MM-DD'


def _header_row(ws, columns: List[str], row: int = 1) -> None:
    for col_idx, title in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=col_idx, value=title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _auto_width(ws, min_width: int = 12, max_width: int = 40) -> None:
    for col in ws.columns:
        length = max(
            len(str(cell.value or "")) for cell in col
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(
            min_width, min(length + 2, max_width)
        )


# ── Sheet 1: Financial Summary ────────────────────────────────────────────────

def _write_summary(ws, transactions: List[Transaction], report_date: date) -> None:
    ws.title = "Financial Summary"
    ws.sheet_view.showGridLines = False

    # Title
    ws.merge_cells("A1:D1")
    title = ws.cell(row=1, column=1, value=f"Financial Summary — {report_date.isoformat()}")
    title.font = Font(bold=True, size=14, name="Calibri")
    title.alignment = Alignment(horizontal="center")

    headers = ["Category", "Credits", "Debits", "Net"]
    _header_row(ws, headers, row=3)

    credits_by_cat: dict[Category, Decimal] = {}
    debits_by_cat: dict[Category, Decimal] = {}
    for t in transactions:
        if t.transaction_type == TransactionType.CREDIT:
            credits_by_cat[t.category] = credits_by_cat.get(t.category, Decimal(0)) + t.amount
        else:
            debits_by_cat[t.category] = debits_by_cat.get(t.category, Decimal(0)) + t.amount

    all_cats = sorted(set(list(credits_by_cat) + list(debits_by_cat)), key=lambda c: c.value)
    row = 4
    total_credits = Decimal(0)
    total_debits = Decimal(0)

    for cat in all_cats:
        cr = credits_by_cat.get(cat, Decimal(0))
        db = debits_by_cat.get(cat, Decimal(0))
        net = cr - db
        total_credits += cr
        total_debits += db

        ws.cell(row=row, column=1, value=cat.value.title()).border = _BORDER
        for col, val in [(2, float(cr)), (3, float(db)), (4, float(net))]:
            c = ws.cell(row=row, column=col, value=val)
            c.number_format = _CURRENCY_FMT
            c.border = _BORDER
            if col == 4:
                c.font = Font(bold=True, color="375623" if net >= 0 else "9C0006", name="Calibri")
        row += 1

    # Totals row
    for col, val, label in [
        (1, "TOTAL", None),
        (2, float(total_credits), None),
        (3, float(total_debits), None),
        (4, float(total_credits - total_debits), None),
    ]:
        c = ws.cell(row=row, column=col, value=val if col != 1 else "TOTAL")
        c.font = _BOLD_FONT
        c.fill = _TOTAL_FILL
        c.border = _BORDER
        if col in (2, 3, 4):
            c.number_format = _CURRENCY_FMT

    _auto_width(ws)
    ws.row_dimensions[3].height = 20


# ── Sheet 2: Transaction Detail ───────────────────────────────────────────────

def _write_transactions(ws, transactions: List[Transaction]) -> None:
    ws.title = "Transaction Detail"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    headers = ["Date", "Type", "Category", "Description", "Vendor", "Reference", "Amount"]
    _header_row(ws, headers)

    sorted_txns = sorted(transactions, key=lambda t: t.date)
    for row_idx, t in enumerate(sorted_txns, start=2):
        fill = _CREDIT_FILL if t.transaction_type == TransactionType.CREDIT else _DEBIT_FILL
        values = [
            t.date,
            t.transaction_type.value,
            t.category.value,
            t.description,
            t.vendor or "",
            t.reference or "",
            float(t.amount),
        ]
        for col_idx, val in enumerate(values, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.fill = fill
            c.border = _BORDER
            c.font = _NORMAL_FONT
            if col_idx == 1:
                c.number_format = _DATE_FMT
            if col_idx == 7:
                c.number_format = _CURRENCY_FMT

    _auto_width(ws)
    ws.row_dimensions[1].height = 20


# ── Sheet 3: AR Aging ─────────────────────────────────────────────────────────

_AGING_BUCKETS = [
    ("Current (not due)", lambda days: days <= 0),
    ("1–30 days",         lambda days: 1 <= days <= 30),
    ("31–60 days",        lambda days: 31 <= days <= 60),
    ("61–90 days",        lambda days: 61 <= days <= 90),
    ("90+ days",          lambda days: days > 90),
]


def _write_ar_aging(ws, invoices: List[Invoice], as_of: date) -> None:
    ws.title = "AR Aging"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:G1")
    title = ws.cell(row=1, column=1, value=f"Accounts Receivable Aging — as of {as_of.isoformat()}")
    title.font = Font(bold=True, size=14, name="Calibri")
    title.alignment = Alignment(horizontal="center")

    headers = ["Invoice #", "Vendor", "Issue Date", "Due Date", "Total", "Days Overdue", "Bucket"]
    _header_row(ws, headers, row=3)

    outstanding = [inv for inv in invoices if not inv.paid]
    bucket_totals: dict[str, Decimal] = {b[0]: Decimal(0) for b in _AGING_BUCKETS}

    for row_idx, inv in enumerate(outstanding, start=4):
        days_overdue = (as_of - inv.due_date).days
        bucket_name = next(
            (name for name, check in _AGING_BUCKETS if check(days_overdue)),
            "90+ days",
        )
        bucket_totals[bucket_name] = bucket_totals.get(bucket_name, Decimal(0)) + inv.total

        fill = _OVERDUE_FILL if days_overdue > 60 else (_WARNING_FILL if days_overdue > 0 else None)
        row_vals = [
            inv.invoice_number, inv.vendor,
            inv.issue_date, inv.due_date,
            float(inv.total), days_overdue, bucket_name,
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.border = _BORDER
            c.font = _NORMAL_FONT
            if fill:
                c.fill = fill
            if col_idx in (3, 4):
                c.number_format = _DATE_FMT
            if col_idx == 5:
                c.number_format = _CURRENCY_FMT

    # Bucket summary below the table
    summary_row = len(outstanding) + 6
    ws.cell(row=summary_row, column=1, value="Aging Summary").font = _BOLD_FONT
    summary_row += 1
    _header_row(ws, ["Bucket", "Total Outstanding"], row=summary_row)
    for label, total in bucket_totals.items():
        summary_row += 1
        ws.cell(row=summary_row, column=1, value=label).border = _BORDER
        c = ws.cell(row=summary_row, column=2, value=float(total))
        c.number_format = _CURRENCY_FMT
        c.border = _BORDER

    _auto_width(ws)


# ── Public API ────────────────────────────────────────────────────────────────

def generate_report(
    transactions: List[Transaction],
    invoices: Optional[List[Invoice]] = None,
    output_path: str | Path = "financial_report.xlsx",
    report_date: Optional[date] = None,
) -> Path:
    """Generate a formatted Excel workbook and return the output path."""
    as_of = report_date or date.today()
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    ws_summary = wb.create_sheet("Financial Summary")
    _write_summary(ws_summary, transactions, as_of)

    ws_detail = wb.create_sheet("Transaction Detail")
    _write_transactions(ws_detail, transactions)

    if invoices:
        ws_aging = wb.create_sheet("AR Aging")
        _write_ar_aging(ws_aging, invoices, as_of)

    wb.save(output_path)
    return output_path
